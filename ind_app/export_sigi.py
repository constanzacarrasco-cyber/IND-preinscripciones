"""Exportación de registros ficticios con la plantilla entregada por IND.

No importa información a SIGI y no está probada contra un ambiente real.
Preserva la hoja auxiliar Options y ambas filas de encabezados de datos.
"""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ARCHIVO_PLANTILLA = Path(__file__).resolve().parent.parent / "assets" / "plantilla_sigi_blanca.xlsx"
ENCABEZADOS_SIGI = (
    "RUT O IDENTIFICACION (obligatorio)",
    "NOMBRE",
    "APELLIDO PATERNO",
    "APELLIDO MATERNO",
    "FECHA NACIMIENTO (dd/mm/aaaa)",
    "GÉNERO (M - F)",
    "CATEGORÍA(SUB 10-12-14-17)",
    "DISCIPLINA",
    "ESTABLECIMIENTO",
    "COMUNA",
)

def _es_plantilla_correcta(hoja) -> bool:
    contenido = [str(hoja.cell(2, i).value or "").strip() for i in range(1,11)]
    return contenido == list(ENCABEZADOS_SIGI)

def generar_excel_sigi(registros: list[dict], archivo_plantilla: Path = ARCHIVO_PLANTILLA) -> bytes:
    """Genera un archivo XLSX a partir de una plantilla SIN datos de alumnos.

    registros: lista de dict creados con el formulario de demostración.
    No se efectúan conexiones de red, escrituras persistentes ni importaciones.
    """
    if not registros:
        raise ValueError("No se puede exportar una lista vacía.")
    if len(registros) > 5000:
        raise ValueError("La demostración admite como máximo 5.000 filas.")
    if not archivo_plantilla.is_file():
        raise FileNotFoundError("Falta la plantilla vacía en assets/plantilla_sigi_blanca.xlsx")
    libro = load_workbook(archivo_plantilla)
    if "datos" not in libro.sheetnames or not _es_plantilla_correcta(libro["datos"]):
        raise ValueError("La plantilla SIGI no tiene las columnas esperadas en la fila 2.")
    hoja = libro["datos"]
    for n, registro in enumerate(registros, start=3):
        fecha = registro["fecha_nacimiento"]
        if isinstance(fecha, datetime):
            fecha = fecha.date()
        if not isinstance(fecha, date):
            raise ValueError(f"El registro de la fila {n} no tiene fecha de nacimiento válida.")
        valores = [
            registro["rut"], registro["nombre"],
            registro["apellido_paterno"], registro["apellido_materno"],
            datetime.combine(fecha, datetime.min.time()),
            registro["genero"], registro["categoria"],
            registro["disciplina"], registro["establecimiento"],
            registro["comuna"],
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=n, column=columna, value=valor)
            if columna == 5:
                celda.number_format = "DD/MM/YYYY"
            else:
                # Tratar el texto como texto para evitar inyección de fórmulas Excel.
                celda.data_type = "s"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
