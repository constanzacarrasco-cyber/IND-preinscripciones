"""Prueba local del generador SIGI, sin conexión a Supabase ni Streamlit."""
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from ind_app.export_sigi import generar_excel_sigi, ENCABEZADOS_SIGI

class TestArchivoSigi(unittest.TestCase):
    def test_encabezados_hojas_y_fecha(self):
        reg=[dict(rut='RUT_FICTICIO',nombre='PRUEBA',apellido_paterno='EJEMPLO',
                  apellido_materno='TEST',fecha_nacimiento=date(2012,5,8),genero='F',
                  categoria='SUB 14',disciplina='Atletismo',
                  establecimiento='ESTABLECIMIENTO FICTICIO',comuna='Coyhaique')]
        raw=generar_excel_sigi(reg)
        libro=load_workbook(BytesIO(raw))
        self.assertIn('datos',libro.sheetnames)
        self.assertIn('Options',libro.sheetnames)
        hoja=libro['datos']
        self.assertEqual(tuple(str(hoja.cell(2,i).value).strip() for i in range(1,11)),ENCABEZADOS_SIGI)
        self.assertEqual(hoja['B3'].value,'PRUEBA')
        self.assertEqual(hoja['E3'].value.date(),date(2012,5,8))
        self.assertEqual(hoja['J3'].value,'Coyhaique')

    def test_lista_vacia_no_exporta(self):
        with self.assertRaises(ValueError):
            generar_excel_sigi([])

if __name__ == '__main__':
    unittest.main()
